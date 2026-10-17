"""Отчёт 1С (xlsx с группировками) → JSON строк с уровнем группировки.

Отчёты «Клиенты» и «Пакеты услуг» 1С:Фитнес выгружаются в Excel деревом:
клуб → статус → (номенклатура →) строка клиента. Уровень строки — outline
level Excel, его читает только полный режим openpyxl (read_only его не
отдаёт), поэтому большой отчёт разбирается пару минут.

    python3 -I dump-report.py клиенты.xlsx clients.json
    python3 -I dump-report.py пакеты.xlsx memberships.json

Дальше — convert.ts (см. его шапку).
"""

import json
import sys

import openpyxl


def main(src: str, out: str) -> None:
    wb = openpyxl.load_workbook(src, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for r in range(1, ws.max_row + 1):
        dim = ws.row_dimensions.get(r)
        level = dim.outlineLevel if dim is not None else 0
        values = []
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            values.append(v.isoformat() if hasattr(v, "isoformat") else v)
        rows.append({"r": r, "lvl": level or 0, "v": values})
    with open(out, "w", encoding="utf-8") as f:
        json.dump(rows, f, ensure_ascii=False)
    print(f"{src}: {len(rows)} строк → {out}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit("usage: dump-report.py <report.xlsx> <out.json>")
    main(sys.argv[1], sys.argv[2])
